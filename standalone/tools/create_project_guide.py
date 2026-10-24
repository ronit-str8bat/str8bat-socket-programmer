"""Create the reviewed local-ap-v9 user guide; no firmware files are changed.

Usage: python create_project_guide.py OUTPUT_DIRECTORY [DEPENDENCY_DIRECTORY]
Dependencies: python-docx, reportlab, PyMuPDF, Pillow.
Content is a documented snapshot, not an automatic extraction of firmware behavior.
"""
import sys
from pathlib import Path
if len(sys.argv) > 2:
    sys.path.insert(0, sys.argv[2])
import html
import json
import math
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image as RLImage
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.pagesizes import A4
import fitz

OUT = Path(sys.argv[1])
OUT.mkdir(parents=True, exist_ok=True)
NAVY = '#12304A'
TEAL = '#007E87'
GRAY = '#536575'
PALE = '#EAF4F5'
RED = '#B43D38'
FONT_DIR = Path('C:/Windows/Fonts')
FONT = FONT_DIR / 'segoeui.ttf'
BOLD = FONT_DIR / 'segoeuib.ttf'
pdfmetrics.registerFont(TTFont('Guide', str(FONT)))
pdfmetrics.registerFont(TTFont('GuideBold', str(BOLD)))
pdfmetrics.registerFontFamily('Guide', normal='Guide', bold='GuideBold', italic='Guide', boldItalic='GuideBold')


class Diagram:
    def __init__(self, height):
        self.im = Image.new('RGB', (1800, height), 'white')
        self.d = ImageDraw.Draw(self.im)

    def text(self, x, y, text, size=29, color=NAVY, bold=False, anchor=None):
        font = ImageFont.truetype(str(BOLD if bold else FONT), size)
        self.d.multiline_text((x, y), text, font=font, fill=color, spacing=9, anchor=anchor)

    def box(self, xy, title, body='', fill=PALE):
        self.d.rounded_rectangle(xy, radius=15, fill=fill, outline=TEAL, width=3)
        x, y, _, _ = xy
        self.text(x+22, y+18, title, 33, bold=True)
        if body:
            self.text(x+22, y+72, body, 29)

    def wire(self, points, color=TEAL, arrow=False):
        self.d.line(points, fill=color, width=5)
        if arrow:
            x, y = points[-1]
            px, py = points[-2]
            a = math.atan2(y-py, x-px)
            self.d.polygon([(x, y), (x-20*math.cos(a-.5), y-20*math.sin(a-.5)),
                            (x-20*math.cos(a+.5), y-20*math.sin(a+.5))], fill=color)

    def dot(self, x, y):
        self.d.ellipse((x-6, y-6, x+6, y+6), fill=TEAL)

    def resistor(self, x1, x2, y, label):
        mid = (x1+x2)//2
        self.wire([(x1, y), (mid-65, y)])
        self.d.rectangle((mid-65, y-17, mid+65, y+17), fill='white', outline=TEAL, width=4)
        self.wire([(mid+65, y), (x2, y)])
        self.text(mid, y-62, label, 26, anchor='mm')

    def save(self, name):
        path = OUT / (name+'.png')
        self.im.save(path, optimize=True)
        return path


d = Diagram(910)
d.box((60, 60, 560, 240), 'Phone / browser', 'Sets the clock over local Wi-Fi\nDownloads and clears CSV logs')
d.box((670, 40, 1220, 420), 'Raspberry Pi Pico 2 W', 'Stored STM32 firmware (.bin)\nSWD programming + verification\nHardware clock and CSV storage\nLocal Wi-Fi access point\nUSB power / diagnostic reports')
d.wire([(560, 150), (670, 150)], arrow=True)
d.text(585, 200, 'Wi-Fi', 27)
d.box((1320, 80, 1740, 290), 'TXS0102', '3.3 V ↔ 1.8 V\nSWCLK + SWDIO\nOE controlled by GP6')
d.wire([(1220, 170), (1320, 170)], arrow=True)
d.box((1300, 520, 1740, 770), 'STM32L476JE', '512 KiB main flash\nPA14 = SWCLK\nPA13 = SWDIO\nNRST = reset')
d.wire([(1530, 290), (1530, 520)], arrow=True)
d.text(1550, 365, 'SWD', 29)
d.wire([(1150, 420), (1150, 635), (1300, 635)], arrow=True)
d.text(950, 560, 'GP1 / NRST', 28)
d.box((60, 510, 550, 720), 'Operator controls', 'START: GP18 to GND\nExternal RGB: GP7 / GP8 / GP9\nOnboard LED: activity / errors')
d.wire([(550, 580), (740, 580), (740, 420)], arrow=True)
d.text(70, 815, 'The phone is optional after clock setup. The Pico programs the STM32 without a computer.', 31, bold=True)
architecture = d.save('01-system-overview')

d = Diagram(1370)
d.text(40, 15, 'REFERENCE CIRCUIT • 1.8 V TARGET • SIGNAL NAMES, NOT IC PACKAGE PIN NUMBERS', 30, bold=True)
d.box((40, 90, 470, 940), 'Pico 2 W', 'Powered through its USB\nconnector (5 V input)')
d.box((660, 230, 1080, 940), 'TXS0102')
d.text(685,390,'B = 3.3 V     A = 1.8 V',26)
d.box((1330, 90, 1760, 940), 'STM32 target board', 'Existing board regulator\nprovides regulated 1.8 V\nand local decoupling')
for y,left,right in [(350,'3V3 OUT · pin 36','VCCB'), (500,'GP2 · pin 4','B1'), (650,'GP3 · pin 5','B2'), (800,'GP6 · pin 9','OE')]:
    d.text(65, y-45, left, 29)
    d.text(680, y-45, right, 29)
    d.wire([(470,y),(660,y)])
for y,left,right in [(350,'VCCA','1.8 V / VDD'), (500,'A1','PA14 / SWCLK'), (650,'A2','PA13 / SWDIO')]:
    d.text(950, y-45, left, 29)
    d.text(1350, y-45, right, 29)
    d.wire([(1080,y),(1330,y)])
d.text(685, 715, 'GP6 HIGH: channels enabled', 25)
d.text(685, 860, 'GND', 29)
d.text(65, 860, 'GND · pin 3', 29)
d.text(1350, 860, 'VSS / GND', 29)
for x in [280,870,1540]:
    d.wire([(x,940),(x,1020)])
    d.dot(x,1020)
d.wire([(120,1020),(1660,1020)])
d.text(620,1045,'COMMON GROUND',29,bold=True)
d.text(80, 1135,'GP1 · pin 2',29,bold=True)
d.wire([(290,1155),(1430,1155)], arrow=True)
d.text(1445,1135,'STM32 NRST',29,bold=True)
d.text(40,1320,'Pico drives LOW to reset; releases the wire to allow NRST HIGH.',28)
d.dot(1170,1155)
d.wire([(1170,1155),(1170,1250)])
d.d.rectangle((1170,1233,1340,1267),fill='white',outline=TEAL,width=4)
d.wire([(1340,1250),(1410,1250)])
d.text(1440,1230,'Target VDD (1.8 V)',27)
d.text(40,1268,'Target-side pull-up: use the board value (10 kΩ is a design example).',25)
swd_diagram = d.save('02-power-swd-reset')

d = Diagram(1310)
d.text(40,15,'POWER AND TRANSLATOR SUPPORT • REFERENCE CONNECTIONS',32,bold=True)
d.box((690,100,1110,420),'LDK120C18R')
d.text(710,190,'IN',29)
d.text(1000,190,'OUT',29)
d.text(710,295,'EN',29)
d.text(845,370,'GND',29)
d.text(50,180,'5 V input',30,bold=True)
d.wire([(245,215),(690,215)])
d.wire([(1110,215),(1710,215)])
d.text(1330,145,'1.8 V → target VDD',28,bold=True)
d.wire([(560,215),(560,320),(690,320)])
d.dot(560,215)
d.text(50,535,'Enable held HIGH independently in this reference; no connection to GP13.',25)
d.wire([(890,420),(890,480)])
d.wire([(245,480),(1710,480)])
d.dot(890,480)
for x,label in [(300,'CIN = 1 µF'),(1250,'COUT = 1 µF')]:
    d.wire([(x,215),(x,325)])
    d.wire([(x-35,325),(x+35,325)])
    d.wire([(x-35,345),(x+35,345)])
    d.wire([(x,345),(x,480)])
    d.dot(x,215); d.dot(x,480)
    d.text(x+45,370,label,25)
d.text(1510,500,'Common GND',27)
for y,left,right in [(660,'Pico 3V3 OUT','TXS0102 VCCB'),(920,'Target 1.8 V','TXS0102 VCCA')]:
    d.text(50,y-23,left,29,bold=True)
    d.wire([(350,y),(1440,y)])
    d.text(1455,y-23,right,28,bold=True)
    d.dot(1000,y)
    d.wire([(1000,y),(1000,y+65)])
    d.wire([(965,y+65),(1035,y+65)])
    d.wire([(965,y+85),(1035,y+85)])
    d.wire([(1000,y+85),(1000,y+155)])
    d.text(1050,y+80,'100 nF',28)
    d.text(890,y+162,'GND',28)
d.text(50,1150,'GP6 (header 9)',29,bold=True)
d.wire([(350,1175),(1440,1175)])
d.text(1455,1150,'TXS0102 OE',28,bold=True)
d.dot(800,1175)
d.wire([(800,1175),(800,1250)])
d.resistor(800,1160,1250,'OE pull-down resistor')
d.text(1190,1230,'GND',28)
power_detail = d.save('02b-power-support')

d = Diagram(1130)
d.text(40,20,'START BUTTON AND EXTERNAL RGB • CURRENT GPIO ASSIGNMENT',32,bold=True)
d.box((40,90,540,320),'START input','GP18 = Pico header pin 24\nInternal pull-up is enabled\nReleased = HIGH; pressed = LOW')
d.wire([(540,200),(800,200)])
d.d.ellipse((794,194,806,206),fill=TEAL)
d.d.ellipse((984,194,996,206),fill=TEAL)
d.wire([(800,200),(970,145)])
d.wire([(990,200),(1300,200)])
d.text(825,235,'Normally open button',28)
d.text(1320,180,'GND',32,bold=True)
d.text(40,370,'OPTION A — COMMON-CATHODE RGB, DIRECT GPIO DRIVE',30,bold=True)
for y,gpio,pin,label,col in [(500,'GP7','10','RED',RED),(680,'GP8','11','GREEN','#198448'),(860,'GP9','12','BLUE','#2463B8')]:
    d.text(60,y-22,f'{gpio} / pin {pin}',30,bold=True)
    d.resistor(320,830,y,'Individual resistor R')
    d.text(890,y-57,label,27,color=col,bold=True)
    d.wire([(830,y),(890,y)])
    d.d.polygon([(890,y-28),(890,y+28),(955,y)],outline=col,width=4)
    d.wire([(963,y-33),(963,y+33)],color=col)
    d.wire([(963,y),(1340,y)])
    d.dot(1340,y)
    d.text(850,y+33,'Anode',25)
d.wire([(1340,500),(1340,960),(1600,960)])
d.text(1360,660,'Shared\ncathode',29)
d.text(1620,940,'GND',30,bold=True)
d.text(45,1020,'Each color needs its own resistor. Select it from LED forward voltage and GPIO current limits.',27)
d.text(45,1065,'R ≈ (3.3 V − LED forward voltage) / chosen LED current. HIGH on the GPIO turns that color ON.',27)
led_diagram = d.save('03-button-and-rgb')

d = Diagram(700)
d.text(40,15,'OPTION B — COMMON-ANODE RGB WITH NPN DRIVERS (ONE COLOR SHOWN)',31,bold=True)
d.text(70,180,'LED supply',30,bold=True)
d.wire([(275,200),(350,200)])
d.resistor(350,650,200,'Current limiter')
d.wire([(650,200),(780,200)])
d.d.polygon([(780,170),(780,230),(845,200)],outline=RED,width=4)
d.wire([(850,166),(850,234)],color=RED)
d.wire([(850,200),(1250,200),(1250,340)])
d.text(745,110,'LED anode → cathode',27)
d.d.ellipse((1120,330,1320,530),outline=TEAL,width=3)
d.wire([(1170,370),(1170,490)])
d.wire([(1170,390),(1250,340)])
d.wire([(1170,470),(1250,525)],arrow=True)
d.wire([(1250,525),(1250,580),(1460,580)])
d.text(1490,560,'GND',30,bold=True)
d.text(70,410,'GP7 / GP8 / GP9',29,bold=True)
d.resistor(370,950,430,'Base resistor / built-in bias')
d.wire([(950,430),(1170,430)])
d.text(1360,365,'NPN transistor\nGPIO HIGH → ON',29)
d.text(60,635,'Repeat for all three colors. The transistor separates the GPIO from the LED supply.',29)
driver_diagram = d.save('04-rgb-driver-alternative')

pages = []
def page(title, subtitle=''):
    elements = []
    pages.append({'title': title, 'subtitle': subtitle, 'elements': elements})
    return elements
def p(e, text): e.append(('p',text))
def h(e, text): e.append(('h',text))
def note(e, text): e.append(('note',text))
def table(e, headers, rows, widths=None): e.append(('table',headers,rows,widths))
def pic(e, path, caption): e.append(('pic',str(path),caption))
def steps(e, items):
    for i,t in enumerate(items,1): p(e,f'<b>{i}.</b> {t}')

e=page('Str8bat Pico Programmer','Circuit, wiring and operator guide • 02 October 2026')
p(e,'<b>A simple guide to programming an STM32, reading the LEDs, setting the clock and collecting CSV records.</b>')
table(e,['Item','Documented configuration'],[
 ['Programmer hardware','Raspberry Pi Pico 2 W / RP2350'],
 ['Programmer firmware','local-ap-v9 • GP6 translator OE • GP18 START'],
 ['Target','STM32L476JE • 512 KiB main flash'],
 ['Embedded target image','str8bat_factory_v2.0_user_v2.0.3.bin • v2.0.3'],
 ['Operator interface','Physical START button, RGB LED, onboard LED and local Wi-Fi webpage'],
 ['After successful programming','Current build leaves the STM32 CPU halted']],[.30,.70])
h(e,'What this project does')
p(e,'A debug probe is a device that communicates with a microcontroller through its debug pins. This project uses that connection to make a standalone programmer: the Pico already holds the STM32 firmware, so a computer is not needed for each programming operation.')
p(e,'The Pico connects through SWD, checks the target, erases main flash, writes the stored image and reads it back to verify it. It records each attempt in a CSV file. A phone connects directly to the Pico to set the time and collect the records.')
h(e,'How to use this guide')
table(e,['Topic','Section'],[['First use and daily operation','1'],['System, circuit and pin connections','2–5'],['External and onboard LEDs','6–7'],['Wi-Fi, webpage and clock','8–10'],['CSV records','11'],['Flash behavior, faults and maintenance','12–14']],[.76,.24])
note(e,'Scope: behavior is checked against the current source and local-hotspot build settings. The circuit drawings are detailed reference connections, not a verified PCB manufacturing schematic. The source itself marks the fixture pin map as provisional; check the actual board before wiring.')

e=page('1. Quick start','One clock setup per Pico power session')
steps(e,[
 'With power off, connect the target using sections 3–5. Use GP18 for START and GP6 for the SWD translator enable. Make sure Pico and STM32 share GND.',
 'Power the Pico and the STM32 board. The target supply must already be enabled. Before the first attempt, the external LED normally alternates green and red while no phone is connected.',
 'On the phone, join <b>Str8bat_Pico</b> using password <b>str8bat@pico</b>. Choose to stay connected if the phone says there is no Internet.',
 'Open the Wi-Fi sign-in page, or browse to <b>http://192.168.4.1/</b>. On a fresh Pico boot, opening the page automatically sends the phone time. Wait for the clock-ready message.',
 'If needed, tap <b>Set clock from this phone</b>. Check that the phone date and time are correct. You may now disconnect the phone; leave the Pico powered.',
 'Press and release START: connect <b>GP18 to GND</b> using the button. Hold long enough for the 100 ms debounce. The external LED turns blue while programming.',
 'Wait for <b>solid green</b> for success, or <b>solid red</b> for failure. If red, count the onboard LED flashes between the long gaps and use section 7.',
 'For another target, safely change the target connections and supply, then press START again. Release the button between attempts. Reconnect the phone later to download the CSV.'
])
note(e,'For correct CSV timestamps: connect the phone and open the webpage before the first programming attempt after each Pico restart. Wi-Fi connection by itself does not set the clock.')
h(e,'What solid green means')
p(e,'The image was programmed, verified and its result saved. In this build, the STM32 is left halted, so green does not mean the target application is running. Reset or power-cycle the STM32 when you want it to run. Keeping the Pico powered preserves its clock.')
h(e,'What needs to remain connected?')
p(e,'During programming: target power, common GND, SWCLK, SWDIO and NRST. A phone, Internet connection and PC are not required after the clock is set.')

e=page('2. System overview','The Pico controls programming; the phone manages time and logs')
pic(e,architecture,'Figure 1. Main parts and their roles in the current local-hotspot build.')
table(e,['Part','Purpose'],[
 ['Pico 2 W','Stores the image, runs the programmer, provides Wi-Fi, keeps time and stores logs.'],
 ['STM32 target board','Receives and runs the firmware after reset/release.'],
 ['TXS0102 translator','Converts the two SWD signal voltage levels. It needs both supply rails; it does not generate either supply. [1]'],
 ['1.8 V regulator / target supply','Provides target power. The earlier schematic names an LDK120C18R regulator for this rail.'],
 ['START and LEDs','Start a job and show ready, busy, success or error.'],
 ['Phone browser','Sets time and downloads/clears CSV logs over the Pico hotspot.']],[.29,.71])
p(e,'This is the standalone local-hotspot application in <b>standalone/src/main.c</b>. The repository also contains the original USB debug-probe application in <b>src/main.c</b>; its behavior is not the operator flow described here.')

e=page('3. Power, SWD and reset circuit','Reference for the 1.8 V target connection discussed in this project')
pic(e,swd_diagram,'Figure 2. Complete programmer-to-target signal connections. A/B channel names are used because TXS0102 package pin numbers vary. [1]')
p(e,'Fit local bypass capacitors close to translator supply pins (100 nF per rail is a reference choice). Provide an OE pull-down so the channels remain off before Pico firmware starts. Choose its value for the actual driver and board. [1]')
p(e,'Use the existing target board regulator, its required capacitors and all STM32 power/ground connections. The drawing treats the assembled STM32 board as one block; it does not replace the MCU package-specific power, reset and decoupling schematic. [3]')
note(e,'Current code does not drive GP13 to enable the regulators. If the PCB has regulator EN connected to GP13, arrange a valid independent enable or revise that design before use. GP4 is unused in this firmware.')

e=page('3. Power circuit details','Regulator, bypass capacitors and OE default state')
pic(e,power_detail,'Figure 2A. Support circuit for Figure 2. Component values are reference choices to check against the assembled board.')
p(e,'The LDK120C18R is a regulator for the 1.8 V rail. Its enable must be HIGH for the output to operate. The datasheet uses 1 µF input/output capacitors; check the exact part, load, thermal dissipation and capacitor requirements. [4]')
p(e,'Figure 2A shows an independently enabled regulator. Do not tie 5 V to a net that is still connected to GP13 or any other Pico GPIO. Use the target board’s intended power-enable arrangement; the current firmware does not control GP13.')
p(e,'TXS0102 VCCA and VCCB are power inputs. It translates signal voltages; it does not produce a 1.8 V or 3.3 V supply. Place bypass capacitors close to the IC, with an OE pull-down for a defined off state before firmware starts. [1]')
note(e,'Use the Pico USB connector for Pico power in this reference. If reusing the earlier PCB, verify J5’s 5 V power routing, regulator part numbers and enables. The supplied drawing labeled one regulator output 3.3 V while showing a 1.8 V part number; that needs checking against the fitted part.')

e=page('4. Exact connection table','GP numbers are GPIO names, not Pico header pin numbers')
table(e,['Pico signal','Header pin','Connect to'],[
 ['GP2 / SWCLK','4','Translator B1 → A1 → STM32 PA14 (SWCLK)'],
 ['GP3 / SWDIO','5','Translator B2 ↔ A2 ↔ STM32 PA13 (SWDIO)'],
 ['GP6 / SWD OE','9','OE of the SWCLK/SWDIO TXS0102'],
 ['GP1 / NRST','2','STM32 NRST; target-side pull-up to target VDD'],
 ['GND','3, or another GND','Target GND and translator GND'],
 ['3V3 OUT','36','Translator VCCB (Pico side)'],
 ['External target 1.8 V','Not a Pico GPIO','Target VDD and translator VCCA (target side)'],
 ['GP18 / START','24','Normally open push button to GND'],
 ['GP7 / red','10','Red LED channel circuit'],
 ['GP8 / green','11','Green LED channel circuit'],
 ['GP9 / blue','12','Blue LED channel circuit']],[.25,.15,.60])
p(e,'The Pico header positions follow the Pico 2 W pinout. [2] Keep SWD wires short and connect only one active programmer to the target SWD pins at a time.')
h(e,'How GP6 enables the translator')
p(e,'GP6 is LOW at startup and while idle. The code sets it HIGH before SWD communication, waits 2 microseconds, and keeps it HIGH while sending and receiving. It sets GP6 LOW when releasing the SWD connection. Reconnection may temporarily disable and re-enable it.')
h(e,'How reset works')
p(e,'NRST is active LOW. GP1 pulls the wire LOW when the target must be reset. To release reset, GP1 becomes an input; the target-side pull-up brings NRST HIGH. The Pico does not drive NRST HIGH. This reset wire is separate from the two SWD translator channels.')
h(e,'Direct wiring and voltage levels')
p(e,'For a compatible 3.3 V target, the direct signal route is GP2 → PA14 and GP3 ↔ PA13, plus GP1 → NRST and common GND. For the reported 1.8 V target, use voltage translation. A direct connection is not electrically equivalent: the Pico uses 3.3 V logic and a 1.8 V response is not guaranteed to meet its input-HIGH requirement. [2]')
note(e,'The old schematic and earlier conversation used different OE assignments. For this v9 guide, GP6 is the SWD translator OE and GP18 is START. Do not connect GP6 only to a reset translator and expect it to enable SWD.')

e=page('5. START and external LED circuit','Choose the LED circuit that matches the actual assembly')
pic(e,led_diagram,'Figure 3. Reference button and common-cathode RGB circuit matching the code comments.')
pic(e,driver_diagram,'Figure 4. Alternative driver circuit for a common-anode RGB, such as the earlier transistor-driven schematic.')
p(e,'The code is active HIGH: GP7, GP8 or GP9 HIGH turns that color on. For direct GPIO drive use a common-cathode RGB and one resistor per color. For a common-anode LED, use appropriate transistor drivers; do not connect a 5 V LED supply directly to a Pico GPIO.')
p(e,'The earlier schematic uses transistor drivers and a shared LED resistor. Confirm its fitted parts and current limits against the PCB; Figure 4 shows the driver principle, not a claim that those exact parts were verified. Button debounce is 100 ms, and holding START does not repeatedly program targets.')

e=page('6. External RGB LED patterns','The external LED gives the overall result')
table(e,['LED pattern','Meaning','What to do'],[
 ['Off briefly at boot','Initial image check / startup.','Wait for startup to finish.'],
 ['Alternating green and red','Ready before the first attempt; no Wi-Fi client is connected.','Connect phone and set clock if this is a fresh Pico boot.'],
 ['Blinking green','Ready before the first attempt; a Wi-Fi client is connected.','Open the webpage and confirm the clock is ready.'],
 ['Solid blue','Programming is in progress.','Keep power and wiring connected.'],
 ['Solid green','Programming, readback verification and log save succeeded.','Target is halted in this build; reset target when ready to run it.'],
 ['Solid red','Attempt failed, or its log could not be saved.','Count the onboard error flashes and read the USB report.'],
 ['Blinking red at startup','Startup/image/radio/storage initialization problem. START is blocked.','Check the build and diagnostic report.']],[.25,.39,.36])
h(e,'Blink timing')
p(e,'The external ready/error patterns change state every 250 ms. Alternating green/red means 250 ms green, then 250 ms red. Blinking green or startup red means 250 ms on, then 250 ms off.')
h(e,'The result stays visible')
p(e,'After the first programming attempt, solid green or red stays displayed until the next attempt. Connecting or disconnecting the phone does not replace the last result with a Wi-Fi pattern.')
note(e,'Blinking green confirms a Wi-Fi client is connected. It does not confirm that the webpage opened or that the clock was set. Check the webpage clock message before starting work.')
h(e,'If the color is wrong')
p(e,'Check the physical RGB type, the GP7/GP8/GP9 connections, the resistor/driver circuit and common GND. A common-anode LED without the intended transistor drivers needs different electrical control from the direct common-cathode circuit.')

e=page('7. Onboard LED patterns','Count flashes between the long gaps')
p(e,'During normal idle and after success, the onboard LED is a heartbeat: approximately 0.5 second ON and 0.5 second OFF. During programming it is faster: 0.125 second ON and 0.125 second OFF.')
p(e,'For an error, each counted flash is 250 ms ON and 250 ms OFF, followed by about a 2-second gap before the count repeats. Count one complete group; do not count the normal heartbeat as an error code.')
table(e,['Flashes','Code meaning','First check'],[
 ['1','Image invalid / unsupported','Correct embedded BIN and a valid UF2 build.'],
 ['2','SWD communication failed','Target power, GND, SWD wires, GP6 OE and reset connection.'],
 ['3','Unexpected target / memory mapping','Correct STM32 device and 512 KiB flash size.'],
 ['4','Protection / option validation failed','Read diagnostic OPTIONS lines; check protection and boot options.'],
 ['5','Reset or halt failed','NRST connection, target reset behavior and supply.'],
 ['6','Operation timed out','Diagnostic stage, target power and communication.'],
 ['7','Flash erase/program error','Flash status in the report and target conditions.'],
 ['8','Readback verification failed','Compare failure address; check supply and signal quality.'],
 ['9','Target clock setup failed','Target clock/power setup details in the report.'],
 ['10','Local log/storage problem','Download and clear saved logs if full; check storage report.'],
 ['11','Clock required but not set','Set phone time through webpage. Not enabled in the current configuration.']],[.11,.39,.50])
note(e,'Two flashes identify a communication failure, not a confirmed BFB2 fault. Four flashes also do not prove that RDP is the cause. Use the report address, stage and OPTIONS fields to identify where the job stopped.')
p(e,'If radio initialization itself fails, the onboard LED may not be available because it is controlled through the Pico W wireless chip. Use the external startup indication and USB report.')

e=page('8. Connect to Wi-Fi and open the page','No router or Internet connection is needed')
table(e,['Setting','Value'],[['Wi-Fi name (SSID)','Str8bat_Pico'],['Password','str8bat@pico'],['Browser address','http://192.168.4.1/'],['Security','WPA2 password-protected local access point']],[.36,.64])
steps(e,[
 'Power the Pico and wait for its startup checks.',
 'Open the phone Wi-Fi settings and select <b>Str8bat_Pico</b>. Enter the password above.',
 'Accept <b>Stay connected</b> or <b>Use without Internet</b> if asked. The hotspot provides a local page, not Internet access.',
 'Tap the Wi-Fi sign-in notification if it appears. If it does not, open Chrome or Safari and enter <b>http://192.168.4.1/</b>. Use HTTP, not HTTPS.',
 'Wait for the page headed <b>STM32 programming logs</b> and the clock-ready message. On a fresh Pico boot, the page sets the clock automatically from the phone.'
])
h(e,'How automatic redirection works')
p(e,'The Pico gives the phone a local network address using DHCP. Its DNS service answers local name requests with the Pico address. When the phone makes an HTTP connectivity check, the Pico returns an HTTP redirect to its own webpage.')
p(e,'The phone may open a small sign-in window or only show a notification. This depends on the phone operating system. The Pico does not redirect encrypted HTTPS pages, so opening the local HTTP address manually is the reliable fallback.')
h(e,'If the page does not open')
p(e,'Confirm the phone is still on Str8bat_Pico. If mobile data, a VPN or automatic network switching sends traffic elsewhere, temporarily disable that feature and try the local address again. For CSV downloads, use the full browser if the sign-in window cannot save files.')
note(e,'The values above are the repository local-hotspot defaults. A deployment can override them in wifi_config.local.h. This profile does not join a factory router and does not need cloud login or NTP time.')

e=page('9. What the webpage can do','Page heading: STM32 programming logs')
table(e,['Control / message','Purpose'],[
 ['Set clock from this phone','Sends the phone date/time to the Pico. Use after a Pico restart or to correct its clock while idle.'],
 ['Download CSV','Downloads programming_logs.csv containing saved programming records.'],
 ['Clear downloaded logs','Clears the current saved records only after a matching completed download. Confirm the file is saved first.'],
 ['Archive links, if present','Download records retained from older firmware/log formats. An absent archive may report unavailable.'],
 ['Clock-ready status','Confirms that the Pico has a valid time reference. It is not a programming progress display.']],[.34,.66])
h(e,'Set the clock')
p(e,'On page load, the browser checks whether the Pico clock is set. If it is not, it automatically sends the current phone time. You can also press the clock button. The success message is <b>Clock set to IST (UTC+05:30). Ready.</b> If already set, the page reports <b>Hardware clock running in IST. Ready.</b>')
h(e,'Download and clear records')
steps(e,[
 'Wait until programming finishes, then tap <b>Download CSV</b>.',
 'Save <b>programming_logs.csv</b> and open it from the phone Downloads/Files app to confirm it is present.',
 'Tap <b>Clear downloaded logs</b> only after confirming the file was saved. Accept the confirmation prompt.',
 'If more records were added after the download, download the new CSV again before clearing. The Pico refuses to erase records that were not in that download.'
])
note(e,'The Pico checks completed network delivery and that the log has not changed. It cannot prove that the phone permanently saved the file. The operator must confirm the saved copy before clearing.')
p(e,'Clearing removes current CSV records, leaving the header. It does not erase the target firmware, reset the clock or remove old archive files. There is no automatic clearing after download.')
p(e,'The page has no programming START button and no firmware upload control in this build. Start programming with GP18. Clock updates and log operations are refused while a job is busy; pressing START can interrupt a download, so download again afterwards.')

e=page('10. How the clock works','Set it once, then disconnect the phone while the Pico stays powered')
steps(e,[
 'The phone browser reads its own current time and sends it as Unix seconds: seconds counted from 1 January 1970 in UTC.',
 'The Pico stores that reference in the RP2350 hardware millisecond timer (POWMAN), using a crystal-derived 1 kHz tick.',
 'That hardware timer continues counting while the Pico runs. Wi-Fi and the browser are not needed to advance it.',
 'For CSV text, the firmware converts UTC to India Standard Time by adding 5 hours 30 minutes.'
])
note(e,'Current UTC time = phone UTC time at setup + elapsed hardware-timer time.\nCSV time in IST = current UTC time + 05:30.')
h(e,'Simple example')
p(e,'Suppose the phone sets the Pico clock at <b>10:00:00 IST</b>. You disconnect the phone. A programming attempt starts 12 minutes 15 seconds later. Its CSV date/time is <b>10:12:15 IST</b>, even though the phone is no longer connected.')
table(e,['Situation','Clock / CSV behavior'],[
 ['Phone disconnects or browser closes','Clock continues while Pico stays powered.'],
 ['Only STM32 is reset or power-cycled','Pico clock continues if Pico power is unchanged.'],
 ['Pico loses power or its application restarts','Open the webpage again to set time for the new session.'],
 ['Programming begins before clock setup','Allowed in this build, but Date&Time(IST) is UNSET.'],
 ['Clock is set after some UNSET records','Only later attempts get a valid timestamp; old rows are not changed.'],
 ['Phone date/time is wrong','Pico receives the wrong reference. Correct the phone and set the clock again.']],[.45,.55])
h(e,'Timestamp versus programming duration')
p(e,'The CSV date/time is captured at the <b>start</b> of the attempt. The Elapsed Time(sec) field uses a separate monotonic duration measurement, so it measures how long programming took; it is not the clock time.')
p(e,'There is no battery-backed time retention or Internet time correction in this profile. Use automatic date/time on the phone. The clock can drift with the oscillator; this setup establishes the time reference, not a precision calibration.')

e=page('11. CSV records and storage','A record of the target, time and programming result')
p(e,'The Pico stores the CSV in its own flash filesystem. Records remain through normal power cycles. They are separate from the STM32 flash and the embedded firmware image.')
table(e,['Column','Meaning'],[
 ['Target UID','STM32 96-bit unique identifier as 24 hexadecimal digits; UNAVAILABLE if it could not be read.'],
 ['Date&Time(IST)','Attempt start time as YYYY-MM-DD HH:MM:SS, in IST; UNSET before clock setup.'],
 ['Elapsed Time(sec)','Programmer duration in seconds, with three decimal places.'],
 ['firmware_version','Build label for the embedded target image; current value v2.0.3.'],
 ['Binary Size(Bytes)','Size of the embedded target image; current value 524288.'],
 ['status','SUCCESS or FAIL. For detailed failure reasons, read the USB report.']],[.32,.68])
h(e,'Example record — illustrative values')
table(e,['Field','Example'],[['Target UID','00112233445566778899AABB'],['Date&Time(IST)','2026-10-02 10:12:15'],['Elapsed Time(sec)','8.450'],['firmware_version','v2.0.3'],['Binary Size(Bytes)','524288'],['status','SUCCESS']],[.40,.60])
p(e,'The example is not a measured programming time or a real target identifier. Open CSV files in Excel, a spreadsheet app or a text editor. Import the UID column as text to preserve leading zeros.')
h(e,'Storage full or interrupted operation')
p(e,'The current CSV is limited to 256 KiB. If the log cannot accept a new job, programming is blocked before erasing the target and the onboard LED shows 10 flashes. Download the records, confirm they are saved and clear them to make space.')
p(e,'A persistent pending-job marker helps record an interrupted attempt as FAIL after reboot. For an interruption where target/time details could not be recovered, fields can show UNAVAILABLE or UNSET, and 0.000 means duration unknown. A power failure can still leave the target only partly programmed.')
note(e,'A logging failure after successful target verification can still produce a red result and 10 flashes. In that case, inspect the report before assuming the STM32 flash was not written.')

e=page('12. What happens inside a programming job','Image checks, BFB2 handling, erase, write and readback')
steps(e,[
 'Validate the embedded image and reserve a log record. Enable SWD, connect to the target and identify the device.',
 'Check the flash size, target state and protection/option settings. If BFB2 is set, clear only BFB2, reload the option bytes and verify the resulting state.',
 'Erase all target main flash. The programmer does not keep old calibration or configuration data unless it is part of the supplied image.',
 'Write the embedded image into target flash. The code uses a RAM loader and fast programming where applicable.',
 'Read back and verify the image. Record the result. In the current build, keep the target CPU halted after success and release the SWD lines.'
])
h(e,'Flash address range for this 512 KiB image')
table(e,['Region','Address range (inclusive)','Size'],[
 ['Full image / main flash','0x08000000 – 0x0807FFFF','512 KiB'],
 ['Bank 1, in dual-bank mode','0x08000000 – 0x0803FFFF','256 KiB'],
 ['Bank 2, in dual-bank mode','0x08040000 – 0x0807FFFF','256 KiB']],[.37,.43,.20])
p(e,'These bank labels assume dual-bank mode with the intended normal mapping. An address increases by one byte; a 32-bit word occupies four consecutive byte addresses. With a 512 KiB image, the programmed address range is not limited to Bank 1.')
h(e,'BFB2 does not indicate whether Bank 2 contains data')
p(e,'BFB2 is a boot-selection option. Writing data at 0x08040000 does not automatically require this option to be checked. The current programmer clears BFB2 only when needed and verifies it after option reload. It preserves the other option settings and does not deliberately remove RDP, WRP or PCROP protection.')
p(e,'A setting that returns after the target application runs needs investigation of that application and reliable option-byte readback. A webpage checkbox or a failed read alone is not proof that the Pico changed every option byte.')
note(e,'The option-handling code uses a 4 MHz MSI clock for reload and later configures the programming clock. Do not remove these steps or erase all option bytes to work around a failure. The USB OPTIONS report distinguishes an attempted clear from a verified clear.')

e=page('13. Troubleshooting and diagnostic reports','Use the reported stage rather than guessing from the LED alone')
table(e,['Symptom','Next action'],[
 ['Red and two flashes','Check power, common GND, NRST, SWD routing and translator enable. Read the latest report.'],
 ['Red and four flashes','Read the option/protection details. Do not blindly change all CubeProgrammer checkboxes.'],
 ['Red and six flashes','A step timed out. Capture the stage/address and check stable power and wiring.'],
 ['Green but target does not run','Expected in the current halt-after-verify build. Reset/power-cycle STM32 to run.'],
 ['CSV time is UNSET','Open the webpage and set clock before subsequent attempts.'],
 ['Clear logs is refused','Finish the download, save it, and re-download if any new records were added.'],
 ['Hotspot connected but page absent','Open http://192.168.4.1/ manually in the full browser.'],
 ['Target supply missing','Check regulator enable. Current firmware does not turn on GP13.']],[.39,.61])
h(e,'Read the USB report')
p(e,'Connect the Pico USB to the PC and run this command from the debugprobe project directory. It only reads the report; it does not send a flash command.')
p(e,'<b>powershell -ExecutionPolicy Bypass -File .\\read-standalone-report.ps1</b>')
p(e,'Save one complete group of NETWORK, PICO-PROGRAMMER, FAST-FIRST and OPTIONS lines. Repeated groups usually describe the same last attempt, not new attempts.')
table(e,['Report field','Meaning'],[['result','Failure category, corresponding to programmer error codes.'],['stage','0 connect/setup; 1 erase; 2 program; 3 verify; 4 restart/final state.'],['address / detail','Register or memory address and diagnostic value; meaning depends on the failure path.'],['OPTIONS before / after','Option values read by the code. after alone is not proof of successful reload.'],['bfb2_cleared','1 means clearing was verified after option reload.'],['halted','Confirmed final target halt state, not a live status after later resets.']],[.34,.66])

e=page('14. Firmware maintenance and references','Keep the source image, version label and deployed UF2 together')
h(e,'Current named build')
p(e,'<b>build-local-hotspot/standalone_programmer-gp6-oe-local-ap-v9-gp18.uf2</b> targets Pico 2 W, enables GP6 OE, uses GP18 START and leaves the target halted after verification. The embedded STM32 firmware label is v2.0.3, size 524288 bytes.')
h(e,'Load or rebuild the Pico firmware')
p(e,'To load the UF2, hold Pico BOOTSEL while connecting its USB to a PC, then copy the correct Pico 2 W UF2 to the RP2350 boot drive. Release BOOTSEL after the drive appears. After the application starts, set the clock again.')
p(e,'To rebuild, supply the desired BIN and its version explicitly. The script defaults currently name a different v4.0.2 image. Replace the example BIN path with its real location:')
p(e,'<b>powershell -ExecutionPolicy Bypass -File .\\build-local-hotspot.ps1 -FirmwarePath "D:\\shashank\\Firmware_bin\\str8bat_factory_v2.0_user_v2.0.3.bin" -FirmwareVersion "v2.0.3"</b>')
p(e,'Do not edit generated/firmware_image.c by hand. It is generated from the BIN. The optional -RunTargetAfterVerify build switch changes the final target behavior; it is not enabled in the documented build.')
h(e,'Source files checked for this guide')
table(e,['File under standalone/','Role'],[['include/programmer_config.h; include/wifi_config.h','Pin map, time policy and Wi-Fi defaults.'],['src/main.c; src/swd.c','Button, LEDs, result handling, SWD and OE.'],['src/stm32_programmer.c; include/stm32_programmer.h','Target programming, options and result codes.'],['src/services/local_portal.c; local_clock.c; local_storage.c','Webpage, hardware clock and CSV behavior.']],[.56,.44])
h(e,'Hardware references')
p(e,'<b>[1] Texas Instruments TXS0102 datasheet</b> — supply rails, signal ports, OE behavior and package pinouts.\nhttps://www.ti.com/lit/ds/symlink/txs0102.pdf')
p(e,'<b>[2] Raspberry Pi Pico 2 W datasheet</b> — board header positions and electrical information.\nhttps://datasheets.raspberrypi.com/picow/pico-2-w-datasheet.pdf')
p(e,'<b>[3] ST STM32L476xx datasheet</b> — target pin functions and board power requirements.\nhttps://www.st.com/resource/en/datasheet/stm32l476je.pdf')
p(e,'<b>[4] ST LDK120 datasheet</b> — regulator enable and reference capacitors.\nhttps://www.st.com/resource/en/datasheet/ldk120.pdf')
note(e,'Review basis: source/configuration inspection and document checks. These drawings are not evidence of a new electrical bench test. Confirm the physical PCB, fitted RGB type and target power arrangement against its actual schematic.')


# Both exports are produced from the same structured content.
formal = '--formal' in sys.argv
if formal:
    from formal_guide_content import build_formal_pages
    pages = build_formal_pages(dict(architecture=architecture, swd_diagram=swd_diagram,
        power_detail=power_detail, led_diagram=led_diagram, driver_diagram=driver_diagram))
styles = {
 'p': ParagraphStyle('body',fontName='Guide',fontSize=10,leading=14,spaceAfter=8,textColor=colors.HexColor(NAVY)),
 'h': ParagraphStyle('subheading',fontName='GuideBold',fontSize=12,leading=16,spaceBefore=7,spaceAfter=6,textColor=colors.HexColor(TEAL)),
 'title': ParagraphStyle('title',fontName='GuideBold',fontSize=23,leading=28,spaceAfter=7,textColor=colors.HexColor(NAVY)),
 'subtitle': ParagraphStyle('subtitle',fontName='Guide',fontSize=10,leading=14,spaceAfter=17,textColor=colors.HexColor(GRAY)),
 'cell': ParagraphStyle('cell',fontName='Guide',fontSize=9.1,leading=12.2,textColor=colors.HexColor(NAVY)),
 'header': ParagraphStyle('header',fontName='GuideBold',fontSize=9.1,leading=12.2,textColor=colors.white),
 'caption': ParagraphStyle('caption',fontName='Guide',fontSize=8.4,leading=11,spaceBefore=4,spaceAfter=9,textColor=colors.HexColor(GRAY)),
 'note': ParagraphStyle('note',fontName='Guide',fontSize=9.6,leading=13.3,textColor=colors.HexColor(NAVY)),
}
WIDTH = A4[0]-88
def para(text, style='p'):
    return Paragraph(text.replace('\n','<br/>'),styles[style])

def footer(canvas, doc):
    canvas.setStrokeColor(colors.HexColor('#D7E2E7'))
    canvas.line(44,43,A4[0]-44,43)
    canvas.setFont('Guide',8)
    canvas.setFillColor(colors.HexColor(GRAY))
    canvas.drawString(44,29,'STR8BAT-DP-UM-001  |  Rev. 1.1' if formal else 'STR8BAT  /  Pico Programmer  /  local-ap-v9')
    canvas.drawRightString(A4[0]-44,29,str(doc.page))

story=[]
doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(A4[0]/72)
sec.page_height=Inches(A4[1]/72)
sec.top_margin=Inches(.58)
sec.bottom_margin=Inches(.60)
sec.left_margin=sec.right_margin=Inches(44/72)
sec.header_distance=Inches(.2)
sec.footer_distance=Inches(.25)
for name in ['Normal','Body Text']:
    st=doc.styles[name]
    st.font.name='Segoe UI'
    st.font.size=Pt(10)
    st.font.color.rgb=RGBColor.from_string(NAVY[1:])
    st.paragraph_format.space_after=Pt(8)
    st.paragraph_format.line_spacing=1.08
for name,size,col in [('Title',23,NAVY),('Heading 1',23,NAVY),('Heading 2',12,TEAL)]:
    st=doc.styles[name]; st.font.name='Segoe UI'; st.font.size=Pt(size); st.font.bold=True
    st.font.color.rgb=RGBColor.from_string(col[1:])
    st.paragraph_format.space_after=Pt(6)
    st.paragraph_format.space_before=Pt(6)
foot=sec.footer.paragraphs[0]
foot.text=('STR8BAT-DP-UM-001 | Rev. 1.1                                      Page ' if formal else 'STR8BAT / Pico Programmer / local-ap-v9                                      ')
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); foot._p.append(field)
for run in foot.runs: run.font.size=Pt(8)

def word_text(container, text, style=None):
    import re
    par=container.add_paragraph(style=style)
    bold=False
    for chunk in re.split('(<b>|</b>)',text):
        if chunk=='<b>': bold=True
        elif chunk=='</b>': bold=False
        else:
            run=par.add_run(html.unescape(chunk)); run.bold=bold
    return par

for ix,pg in enumerate(pages):
    if ix:
        story.append(PageBreak()); doc.add_page_break()
    story.extend([para(pg['title'],'title'),para(pg['subtitle'],'subtitle')])
    doc.add_heading(pg['title'],level=1)
    sub=doc.add_paragraph(pg['subtitle']); sub.paragraph_format.space_after=Pt(14)
    for run in sub.runs: run.font.color.rgb=RGBColor.from_string(GRAY[1:])
    for el in pg['elements']:
        kind=el[0]
        if kind in ('p','h'):
            story.append(para(el[1],kind))
            word_text(doc,el[1], 'Heading 2' if kind=='h' else None)
        elif kind=='note':
            box=Table([[para(el[1],'note')]],colWidths=[WIDTH])
            box.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,-1),colors.HexColor(PALE)),('BOX',(0,0),(-1,-1),.6,colors.HexColor('#B5D7D9')),('LEFTPADDING',(0,0),(-1,-1),10),('RIGHTPADDING',(0,0),(-1,-1),10),('TOPPADDING',(0,0),(-1,-1),9),('BOTTOMPADDING',(0,0),(-1,-1),9)]))
            story.extend([box,Spacer(1,9)])
            par=word_text(doc,el[1]); par.paragraph_format.space_before=Pt(5)
            shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),PALE[1:]); par._p.get_or_add_pPr().append(shade)
        elif kind=='table':
            _,heads,rows,widths=el
            fractions=widths or [1/len(heads)]*len(heads)
            data=[[para(html.escape(str(s)),'header') for s in heads]]+[[para(html.escape(str(s)),'cell') for s in row] for row in rows]
            t=Table(data,colWidths=[WIDTH*x for x in fractions],repeatRows=1,hAlign='LEFT')
            t.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor(NAVY)),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F1F5F7')]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),6),('BOTTOMPADDING',(0,0),(-1,-1),6),('LEFTPADDING',(0,0),(-1,-1),7),('RIGHTPADDING',(0,0),(-1,-1),7),('LINEBELOW',(0,0),(-1,-1),.35,colors.HexColor('#D6E1E7'))]))
            story.extend([t,Spacer(1,10)])
            wt=doc.add_table(rows=1,cols=len(heads)); wt.autofit=False
            for i,f in enumerate(fractions): wt.columns[i].width=Inches(WIDTH*f/72)
            for rowix,row in enumerate([heads]+rows):
                cells=wt.rows[0].cells if rowix==0 else wt.add_row().cells
                for j,value in enumerate(row):
                    cell=cells[j]; cell.width=Inches(WIDTH*fractions[j]/72)
                    cp=cell.paragraphs[0]; cp.paragraph_format.space_after=Pt(4); cp.paragraph_format.space_before=Pt(4)
                    cp.paragraph_format.line_spacing=1.0
                    rr=cp.add_run(str(value)); rr.font.size=Pt(9); rr.bold=rowix==0
                    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),NAVY[1:] if rowix==0 else ('F1F5F7' if rowix%2==0 else 'FFFFFF')); cell._tc.get_or_add_tcPr().append(sh)
                    if rowix==0: rr.font.color.rgb=RGBColor(255,255,255)
                trpr=wt.rows[rowix]._tr.get_or_add_trPr(); prevent=OxmlElement('w:cantSplit'); trpr.append(prevent)
                if rowix==0:
                    repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
            after=doc.add_paragraph(); after.paragraph_format.space_after=Pt(0); after.paragraph_format.space_before=Pt(0); after.paragraph_format.line_spacing=Pt(3)
        elif kind=='pic':
            _,path,caption=el
            with Image.open(path) as im: iw,ih=im.size
            story.extend([RLImage(path,width=WIDTH,height=WIDTH*ih/iw),para(caption,'caption')])
            doc.add_picture(path,width=Inches(WIDTH/72))
            cap=doc.add_paragraph(caption)
            for run in cap.runs: run.font.size=Pt(8); run.font.color.rgb=RGBColor.from_string(GRAY[1:])
            cap.paragraph_format.space_after=Pt(6)

stem='Str8bat_Pico_Programmer_Technical_Manual_Rev1_1' if formal else 'Str8bat_Pico_Programmer_Guide_v9'
pdf=OUT/(stem+'.pdf')
word=OUT/(stem+'.docx')
SimpleDocTemplate(str(pdf),pagesize=A4,rightMargin=44,leftMargin=44,topMargin=40,bottomMargin=55,
                  title='Str8bat Pico Programmer — Circuit, Wiring and Operator Guide',author='Str8bat Project Documentation').build(story,onFirstPage=footer,onLaterPages=footer)
doc.core_properties.title='Str8bat Pico Programmer — Technical Description and Operating Procedure' if formal else 'Str8bat Pico Programmer — Circuit, Wiring and Operator Guide'
doc.core_properties.subject='local-ap-v9 wiring, LEDs, Wi-Fi, clock and CSV logging'
doc.core_properties.author='Str8bat Project Documentation'
doc.save(word)
(OUT/'guide-content.json').write_text(json.dumps(pages,ensure_ascii=False,indent=2),encoding='utf-8')

pdfdoc=fitz.open(pdf)
checks={'pdf_pages':len(pdfdoc),'planned_sections':len(pages),'files':[str(pdf),str(word)],'pages':[]}
thumbs=[]
for i,pg in enumerate(pdfdoc):
    text=pg.get_text()
    checks['pages'].append({'page':i+1,'chars':len(text),'first_line':text.splitlines()[1:3]})
    pix=pg.get_pixmap(matrix=fitz.Matrix(.5,.5),alpha=False)
    im=Image.frombytes('RGB',[pix.width,pix.height],pix.samples)
    thumbs.append(im)
cols=3; rows=math.ceil(len(thumbs)/cols)
sheet=Image.new('RGB',(cols*320,rows*450),'#CAD5DE')
for i,im in enumerate(thumbs): sheet.paste(im,(i%cols*320+10,i//cols*450+10))
sheet.save(OUT/'preview-contact-sheet.png')
(OUT/'validation.json').write_text(json.dumps(checks,indent=2),encoding='utf-8')
print(json.dumps(checks,indent=2))
